"""Offline behavioral checks: no real API key, model, or network required."""
import io
import unittest
from pathlib import Path
from unittest.mock import patch, Mock
from docx import Document as WordDocument
from langchain_core.documents import Document
from agent import build_agent, new_run_config, start_run, resume_run
from evaluate import compare_results, load_answer_key, main, default_test_sets
from retrieval import HybridSearch
from tools import Judgment, judge_rule, load_rules, read_document, JUDGE_PROMPT, quote_span
import config
from api_calls import call_with_timeout_retry


class FakeModel:
    def __init__(self, judgments):
        self.judgments = iter(judgments)
        self.calls = 0

    def with_structured_output(self, schema):
        return self

    def invoke(self, messages):
        self.calls += 1
        return next(self.judgments)


class FakeIndex:
    def __init__(self, chunks):
        self.chunks = chunks
        self.queries = []

    def search(self, query):
        self.queries.append(query)
        return self.chunks


def judgment(strong=True, quote='Invoices are payable in 30 days.'):
    return Judgment(verdict='MET' if strong else 'UNCLEAR', quote=quote,
                    chunk_id=0, explanation='Payment terms.', reason='Payment terms.', evidence_strong=strong,
                    relevant_evidence=True,
                    alternate_query='Payment deadline and invoice terms')


class ProjectTests(unittest.TestCase):
    def test_absence_and_related_evidence_verdicts(self):
        absent = Judgment(verdict='UNCLEAR', quote='', chunk_id=-1,
                          explanation='Nothing relevant.', reason='No evidence found in document', evidence_strong=False,
                          relevant_evidence=False, alternate_query='Invoice payment deadline')
        contradiction = judgment().model_copy(update={'verdict': 'NOT MET'})
        invalid = judgment(quote='Invented quote')
        cases = [
            ('absent twice', [absent, absent], 'NOT MET', 2),
            ('related then absent', [judgment(False), absent], 'NOT MET', 2),
            ('absent then related', [absent, judgment(False)], 'UNCLEAR', 2),
            ('absent then fulfilled', [absent, judgment()], 'MET', 2),
            ('absent then contradicted', [absent, contradiction], 'NOT MET', 2),
            ('contradiction', [contradiction], 'NOT MET', 1),
            ('invalid then absent', [invalid, absent], 'NOT MET', 2),
            ('invalid twice', [invalid, invalid], 'UNCLEAR', 2),
        ]
        for name, judgments, expected, attempts in cases:
            with self.subTest(name=name):
                model = FakeModel(judgments)
                graph = build_agent(model, FakeIndex)
                run = new_run_config()
                start_run(graph, run, b'Invoices are payable in 30 days.', 'a.txt', 'Payment due?')
                result = resume_run(graph, run, True)['results'][0]
                self.assertEqual(result['verdict'], expected)
                self.assertEqual(result['search_attempts'], attempts)
                self.assertEqual(model.calls, attempts)
                self.assertTrue(result['reason'])
                if judgments[-1] is absent:
                    self.assertEqual(result['explanation'], 'No evidence found in document')
                    self.assertEqual(result['reason'], 'No evidence found in document')
                    self.assertEqual(result['quote'], '')
                    self.assertEqual(result['location'], '')
                elif name == 'contradiction':
                    self.assertTrue(result['quote'])
                    self.assertTrue(result['location'])

    def test_first_absence_is_provisional_and_irrelevant_chunks_do_not_count(self):
        absent = Judgment(verdict='UNCLEAR', quote='', chunk_id=-1,
                          explanation='Unrelated content.', reason='No evidence found in document', evidence_strong=False,
                          relevant_evidence=False, alternate_query='Governing jurisdiction')
        chunks = read_document(b'Invoices are payable in 30 days.', 'a.txt')
        result = judge_rule(FakeModel([absent]), {'id': '1', 'rule': 'Governing law?'}, chunks)
        self.assertEqual(result['verdict'], 'NOT MET')
        self.assertTrue(result['no_evidence'])
        self.assertFalse(result['strong'])

    def test_governing_law_absent_after_retry(self):
        unrelated = judgment(False)
        absent = unrelated.model_copy(update={
            'verdict': 'UNCLEAR', 'quote': '', 'chunk_id': -1,
            'evidence_strong': False, 'relevant_evidence': False,
            'reason': 'No evidence found in document'})
        graph = build_agent(FakeModel([unrelated, absent]), FakeIndex)
        run = new_run_config()
        start_run(graph, run, b'Invoices are payable in 30 days.', 'a.txt',
                  'id,rule\nR4,The contract must specify governing law.', True)
        result = resume_run(graph, run, True)['results'][0]
        self.assertEqual(result['id'], 'R4')
        self.assertEqual(result['verdict'], 'NOT MET')
        self.assertEqual(result['reason'], 'No evidence found in document')

    def test_evaluation_prints_reason_for_every_verdict(self):
        rules = load_rules(Path('examples/rules.csv').read_text(), True)
        answers = load_answer_key('examples/answers.csv', rules)
        results = [{'id': key, 'verdict': value, 'reason': f'Reason for {key}'}
                   for key, value in answers.items()]
        paused = {'__interrupt__': [type('Pause', (), {'value': {'rules': rules}})()]}
        with patch('sys.argv', ['evaluate.py', '--approve', '--document', 'examples/document.txt',
                                '--rules', 'examples/rules.csv', '--answers', 'examples/answers.csv']), \
             patch('evaluate.build_agent'), patch('evaluate.start_run', return_value=paused), \
             patch('evaluate.resume_run', return_value={'results': results}), \
             patch('sys.stdout', new_callable=io.StringIO) as output:
            main()
        for result in results:
            self.assertIn(f"{result['id']}: {result['verdict']} — {result['reason']}",
                          output.getvalue())

    def test_absence_does_not_depend_on_verdict_chunk_or_strength(self):
        chunks = read_document(b'Invoices are payable in 30 days.', 'a.txt')
        for verdict in ('MET', 'NOT MET', 'UNCLEAR'):
            for quote in ('', '   '):
                with self.subTest(verdict=verdict, quote=quote):
                    absent = judgment().model_copy(update={
                        'verdict': verdict, 'quote': quote, 'chunk_id': 0,
                        'relevant_evidence': False, 'evidence_strong': True})
                    result = judge_rule(FakeModel([absent]),
                                        {'id': 'R4', 'rule': 'Governing law?'}, chunks)
                    self.assertTrue(result['no_evidence'])
                    self.assertEqual(result['verdict'], 'NOT MET')
                    self.assertEqual(result['reason'], 'No evidence found in document')

    def test_default_sets_and_total_with_wrong_reason(self):
        outcomes = []
        for name, document, rules_path, answers_path in default_test_sets():
            self.assertTrue(read_document(document.read_bytes(), document.name))
            rules = load_rules(rules_path.read_text(encoding='utf-8-sig'), True)
            answers = load_answer_key(answers_path, rules)
            results = [{'id': key, 'verdict': verdict, 'reason': 'Test reason'}
                       for key, verdict in answers.items()]
            results[0]['verdict'] = 'UNCLEAR' if results[0]['verdict'] != 'UNCLEAR' else 'MET'
            outcomes.append({'results': results})
        total = sum(len(outcome['results']) for outcome in outcomes)
        paused = {'__interrupt__': [type('Pause', (), {'value': {'rules': []}})()]}
        with patch('sys.argv', ['evaluate.py', '--approve']), \
             patch('evaluate.build_agent'), patch('evaluate.start_run', return_value=paused), \
             patch('evaluate.resume_run', side_effect=outcomes), \
             patch('sys.stdout', new_callable=io.StringIO) as output:
            main()
        text = output.getvalue()
        for (name, *_), outcome in zip(default_test_sets(), outcomes):
            count = len(outcome['results'])
            self.assertIn(f'{name}: {count - 1}/{count} correct', text)
        self.assertEqual(text.count('Reason: Test reason'), 3)
        self.assertIn(f'Total: {total - 3}/{total} correct', text)

    def test_api_failure_is_displayed_and_other_sets_continue(self):
        with patch('sys.argv', ['evaluate.py', '--approve']), \
             patch('evaluate.evaluate_set', side_effect=[TimeoutError('Request timed out'),
                                                         (2, 3), (4, 5)]) as run, \
             patch('sys.stdout', new_callable=io.StringIO) as output, \
             patch('sys.stderr', new_callable=io.StringIO):
            with self.assertRaises(SystemExit) as error:
                main()
        self.assertEqual(error.exception.code, 2)
        self.assertEqual(run.call_count, 3)
        self.assertIn('ERROR (TimeoutError): Request timed out', output.getvalue())
        self.assertIn('Total (completed sets): 6/8 correct', output.getvalue())

    def test_openai_clients_have_timeout_no_retries_and_zero_temperature(self):
        self.assertEqual(config.TEMPERATURE, 0)
        self.assertEqual(config.MAX_RETRIES, 0)
        self.assertEqual(config.REQUEST_TIMEOUT, 60)
        self.assertEqual(config.TIMEOUT_RETRIES, 1)
        graph = build_agent(index_factory=FakeIndex)
        run = new_run_config()
        start_run(graph, run, b'Invoices are payable in 30 days.', 'a.txt', 'Payment?')
        with patch('agent.ChatOpenAI', return_value=FakeModel([judgment()])) as chat:
            resume_run(graph, run, True)
        self.assertEqual(chat.call_args.kwargs['temperature'], 0)
        self.assertEqual(chat.call_args.kwargs['timeout'], 60)
        self.assertEqual(chat.call_args.kwargs['max_retries'], 0)
        chunks = read_document(b'Invoices are payable in 30 days.', 'a.txt')
        with patch('retrieval.OpenAIEmbeddings') as embeddings:
            embeddings.return_value.embed_documents.return_value = [[1, 0]]
            HybridSearch(chunks)
        self.assertEqual(embeddings.call_args.kwargs['request_timeout'], 60)
        self.assertEqual(embeddings.call_args.kwargs['max_retries'], 0)

    def test_approval_prevents_model_and_index_calls(self):
        model = FakeModel([judgment()])
        indexes = []
        def factory(chunks):
            indexes.append(FakeIndex(chunks))
            return indexes[-1]
        graph = build_agent(model, factory)
        run = new_run_config()
        paused = start_run(graph, run, b'Invoices are payable in 30 days.', 'sample.txt',
                           'Invoices must be paid in 30 days.')
        self.assertIn('__interrupt__', paused)
        self.assertEqual(model.calls, 0)
        self.assertEqual(indexes, [])
        result = resume_run(graph, run, True)
        self.assertEqual(result['results'][0]['verdict'], 'MET')
        self.assertEqual(result['results'][0]['search_attempts'], 1)
        self.assertIn('characters 1', result['results'][0]['location'])

    def test_rejection_ends_without_calls(self):
        graph = build_agent(FakeModel([]), lambda _: self.fail('Index built before approval'))
        run = new_run_config()
        start_run(graph, run, b'Hello', 'a.txt', 'A rule')
        result = resume_run(graph, run, False)
        self.assertFalse(result['approved'])
        self.assertEqual(result['results'], [])

    def test_weak_evidence_retries_once_then_stops(self):
        model = FakeModel([judgment(False), judgment(False)])
        indexes = []
        def factory(chunks):
            indexes.append(FakeIndex(chunks))
            return indexes[-1]
        graph = build_agent(model, factory)
        run = new_run_config()
        start_run(graph, run, b'Invoices are payable in 30 days.', 'a.txt', 'Payment due?')
        result = resume_run(graph, run, True)['results'][0]
        self.assertEqual(result['verdict'], 'UNCLEAR')
        self.assertEqual(model.calls, 2)
        self.assertEqual(len(indexes[0].queries), 2)
        self.assertNotEqual(*indexes[0].queries)

    def test_invented_quote_is_removed(self):
        chunks = read_document(b'Invoices are payable in 30 days.', 'a.txt')
        result = judge_rule(FakeModel([judgment(quote='Invented quote')]),
                            {'id': '1', 'rule': 'Payment?'}, chunks)
        self.assertEqual(result['verdict'], 'UNCLEAR')
        self.assertEqual(result['quote'], '')
        self.assertEqual(result['location'], '')

    def test_tolerant_table_quote_returns_original_and_location(self):
        document = WordDocument()
        table = document.add_table(rows=1, cols=3)
        for cell, text in zip(table.rows[0].cells,
                              ['Battery – temperature', '≤ 60 °C', '72 °C']):
            cell.text = text
        output = io.BytesIO()
        document.save(output)
        chunks = read_document(output.getvalue(), 'table.docx')
        copied = 'Battery - temperature <= 60 C 72 C'
        answer = judgment(quote=copied).model_copy(update={'verdict': 'NOT MET'})
        result = judge_rule(FakeModel([answer]),
                            {'id': 'TF4', 'rule': 'Test results must not exceed their limits.'}, chunks)
        self.assertEqual(result['verdict'], 'NOT MET')
        self.assertEqual(result['quote'], chunks[0].page_content)
        self.assertIn(f"characters 1–{len(result['quote'])}", result['location'])
        self.assertIsNotNone(quote_span('≤ 100 µA', '<= 100 μA'))
        self.assertIsNotNone(quote_span('≤ 100 µA', '<= 100 uA'))
        self.assertIsNone(quote_span('≤ 60 °C', '<= 72 C'))
        self.assertIsNone(quote_span('≤ 60 °C', '>= 60 C'))

    def test_placeholder_section_is_not_met(self):
        self.assertIn('"To be completed", "TBD", or is blank', JUDGE_PROMPT)
        for placeholder in ('To be completed.', 'TBD'):
            answer = judgment(quote=placeholder).model_copy(update={
                'verdict': 'NOT MET', 'reason': 'Required section is a placeholder.'})
            graph = build_agent(FakeModel([answer]), FakeIndex)
            run = new_run_config()
            start_run(graph, run, placeholder.encode(), 'a.txt',
                      'Labelling and instructions for use must be provided.')
            result = resume_run(graph, run, True)['results'][0]
            self.assertEqual(result['verdict'], 'NOT MET')
            self.assertEqual(result['quote'], placeholder)

    def test_prohibition_absence_is_met_after_retry_and_checks_all_chunks(self):
        absent = judgment(False, '').model_copy(update={
            'verdict': 'MET', 'relevant_evidence': False})
        class EmptyIndex(FakeIndex):
            def search(self, query):
                return []
        class InspectModel(FakeModel):
            def invoke(self, messages):
                self_outer.assertIn('ordinary device', messages[1][1])
                return super().invoke(messages)
        self_outer = self
        model = InspectModel([absent, absent])
        graph = build_agent(model, EmptyIndex)
        run = new_run_config()
        start_run(graph, run, b'An ordinary device.', 'a.txt',
                  'The document must not claim that the device is "100% safe".')
        result = resume_run(graph, run, True)['results'][0]
        self.assertEqual(result['verdict'], 'MET')
        self.assertEqual(result['reason'], 'No violation found in document')
        self.assertEqual(result['search_attempts'], 2)
        self.assertEqual(result['quote'], '')

    def test_prohibition_violation_is_not_met(self):
        answer = judgment(quote='100% safe').model_copy(update={'verdict': 'NOT MET'})
        graph = build_agent(FakeModel([answer]), FakeIndex)
        run = new_run_config()
        start_run(graph, run, b'This device is 100% safe.', 'a.txt',
                  'The document must not claim that the device is "100% safe".')
        result = resume_run(graph, run, True)['results'][0]
        self.assertEqual(result['verdict'], 'NOT MET')
        self.assertEqual(result['quote'], '100% safe')

    def test_timeout_retry_once_and_other_errors_fail_immediately(self):
        from openai import APITimeoutError
        import httpx
        timeout = APITimeoutError(request=httpx.Request('POST', 'https://example.com'))
        for error in (TimeoutError('timed out'), timeout):
            operation = Mock(side_effect=[error, 'success'])
            self.assertEqual(call_with_timeout_retry(operation), 'success')
            self.assertEqual(operation.call_count, 2)
            operation = Mock(side_effect=error)
            with self.assertRaises(type(error)):
                call_with_timeout_retry(operation)
            self.assertEqual(operation.call_count, 2)
        operation = Mock(side_effect=ValueError('bad request'))
        with self.assertRaisesRegex(ValueError, 'bad request'):
            call_with_timeout_retry(operation)
        self.assertEqual(operation.call_count, 1)

    def test_chat_and_embeddings_retry_timeouts(self):
        model = Mock()
        model.with_structured_output.return_value.invoke.side_effect = [
            TimeoutError('timeout'), judgment()]
        chunks = read_document(b'Invoices are payable in 30 days.', 'a.txt')
        result = judge_rule(model, {'id': '1', 'rule': 'Payment?'}, chunks)
        self.assertEqual(result['verdict'], 'MET')
        embeddings = Mock()
        embeddings.embed_documents.side_effect = [TimeoutError(), [[1, 0]]]
        embeddings.embed_query.side_effect = [TimeoutError(), [1, 0]]
        index = HybridSearch(chunks, embeddings)
        self.assertEqual(index.search('Payment'), chunks)
        self.assertEqual(embeddings.embed_documents.call_count, 2)
        self.assertEqual(embeddings.embed_query.call_count, 2)

    def test_rule_csv_validation(self):
        self.assertEqual(load_rules('\ufeffid,rule\nR1,Check this\n', True)[0]['id'], 'R1')
        for text in ['wrong\nfoo', 'id,rule\n1,a\n1,b', 'id,rule\n1,']:
            with self.assertRaises(ValueError):
                load_rules(text, True)

    def test_docx_tables_are_loaded(self):
        document = WordDocument()
        document.add_paragraph('Start')
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Table evidence'
        output = io.BytesIO()
        document.save(output)
        chunks = read_document(output.getvalue(), 'a.docx')
        self.assertEqual(chunks[1].page_content, 'Table evidence')
        self.assertEqual(chunks[1].metadata['location'], 'block 2')

    def test_text_chunk_offsets_preserve_source(self):
        source = 'A paragraph with exact whitespace.\n\n' * 100
        for chunk in read_document(source.encode(), 'a.txt'):
            start = chunk.metadata['start_index']
            self.assertEqual(source[start:start + len(chunk.page_content)], chunk.page_content)

    def test_pdf_loader(self):
        from pypdf import PdfWriter
        output = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(output)
        with self.assertRaisesRegex(ValueError, 'OCR'):
            read_document(output.getvalue(), 'scan.pdf')

    def test_hybrid_search_uses_both_methods(self):
        class Embeddings:
            def embed_documents(self, texts):
                return [[1, 0], [0, 1], [-1, 0]]
            def embed_query(self, text):
                return [1, 0]
        chunks = [Document(page_content=t) for t in ['semantic evidence', 'keyword', 'other']]
        index = HybridSearch(chunks, Embeddings())
        with patch.object(config, 'TOP_K', 2):
            self.assertEqual(set(c.page_content for c in index.search('keyword')),
                             {'semantic evidence', 'keyword'})

    def test_scoring_and_answer_key(self):
        rules = load_rules(Path('examples/rules.csv').read_text(), True)
        answers = load_answer_key('examples/answers.csv', rules)
        results = [{'id': k, 'verdict': v} for k, v in answers.items()]
        results[0]['verdict'] = 'UNCLEAR'
        score, wrong = compare_results(results, answers)
        self.assertEqual(score, 5)
        self.assertEqual(wrong[0]['id'], 'R1')
        with self.assertRaises(ValueError):
            load_answer_key('examples/answers.csv', rules[:1])

    def test_maximum_rules_finish_even_with_all_retries(self):
        model = FakeModel([judgment(False)] * (config.MAX_RULES * 2))
        graph = build_agent(model, FakeIndex)
        run = new_run_config()
        start_run(graph, run, b'Invoices are payable in 30 days.', 'a.txt',
                  '\n'.join(f'Rule {i}' for i in range(config.MAX_RULES)))
        outcome = resume_run(graph, run, True)
        self.assertEqual(len(outcome['results']), config.MAX_RULES)
        self.assertEqual(model.calls, config.MAX_RULES * 2)


if __name__ == '__main__':
    unittest.main()
